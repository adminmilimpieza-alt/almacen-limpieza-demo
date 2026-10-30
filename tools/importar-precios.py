#!/usr/bin/env python3
"""
importar-precios.py — Convierte la lista de precios del proveedor en el catálogo
semilla de la app (data/catalogo.json y js/catalogo.js).

La app en producción NO usa este archivo para actualizar el negocio: la carga se
hace desde Precios → "Importar la lista completa del proveedor". Este script sólo
deja la semilla correcta para una instalación nueva y permite revisar el catálogo
en el repositorio.

Uso:
    python tools/importar-precios.py "MILIMPIEZA-Lista Precios-Sofware.xlsx"

Estructura de la planilla (hoja "Lista PRECIO"):
    A=Código  B=Descripción  C=Venta  D=Cant./Lt.  E=Costo  F=% Ganancia  H=Costo botella
Cada fila con A="Codigo" abre una sección y el proveedor viene en B:
    "Descripcion: LINEA PILETA ... Proveedor: UNIBLANC"
El código está una sola vez (celda combinada) y las presentaciones siguientes lo heredan.

Regla del cliente: la venta se calcula desde costo y margen, y la botella se suma
a la venta. Por eso el costo que se guarda es venta ÷ (1 + margen/%): así la app
reproduce exactamente el precio de la planilla. El envase queda documentado en
"alerta" para poder auditarlo.
"""
import json
import math
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "js", "catalogo.js")
SALIDA_PUBLICADO = os.path.join(RAIZ, "data", "catalogo.json")

ENVASES = re.compile(r"^(BID[OÓ]N|BOTELLA|ENVASE|PULVERIZADOR|ROCIADOR|ASPERSOR)", re.I)
PRESENTACIONES = [
    (r"\b10\s*LTS?", "10 Lts"), (r"\b5\s*LTS?", "5 Lts"), (r"\b1\s*LT\.?(?![A-Z])", "1 Lt"),
    (r"\b50\s*KGS?", "50 Kg"), (r"\b45\s*KGS?", "45 Kg"), (r"\b1\s*KGS?\.?", "1 Kg"),
    (r"\b1\s*KGR", "1 Kg"), (r"\b500\s*GR", "500 grs"), (r"\b250\s*CC", "250 cc"),
    (r"\b360\s*CC", "360 cc"), (r"\b750\s*ML", "750 ml"),
]


def redondear(valor):
    """Mismo redondeo que la app (Math.round(x*100)/100), sin sorpresas de centavos."""
    return math.floor(valor * 100 + 0.5) / 100


def normaliza(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", texto.upper())


def limpia(valor):
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def presentacion(detalle, cant):
    completo = detalle or ""
    # El tamaño del envase va antes que las dosis o rindes que la planilla pone entre paréntesis.
    fuera = re.sub(r"\([^)]*\)", " ", completo)
    for patron, valor in PRESENTACIONES:
        if re.search(patron, fuera, re.I):
            return valor
    if isinstance(cant, (int, float)) and cant > 0:
        if cant == 1:
            return "1 Lt"
        if cant == 5:
            return "5 Lts"
        if cant == 10:
            return "10 Lts"
        if cant == 0.5:
            return "500 grs"
        if float(cant).is_integer():
            return f"{int(cant)} x Unid"
    for patron, valor in PRESENTACIONES:
        if re.search(patron, completo, re.I):
            return valor
    return "x Unid"


def nombre_corto(detalle, pres):
    base = limpia(detalle)
    parentesis = base.find("(")
    if parentesis > 12:
        base = base[:parentesis]
    elif ". " in base:
        base = base.split(". ")[0]
    base = re.sub(r"[.;,]+$", "", limpia(base))
    final = re.search(r"(\d+)\s*(LTS?|LT|KGR|GR|CC|ML)\s*$", base, re.I)
    if final and pres != "x Unid":
        real = re.match(r"(\d+)", pres)
        if real and real.group(1) != final.group(1):
            base = limpia(base[: final.start()]) + " " + pres
    return base or limpia(detalle)


def proveedor_de(cabecera):
    etiqueta = re.search(r"Proveedor:\s*([^|]+)$", cabecera, re.I)
    if etiqueta:
        return limpia(etiqueta.group(1))
    solo = re.sub(r"^Descripcion:\s*", "", cabecera, flags=re.I).strip()
    return solo.upper() if re.search(r"GEMORA|DON EMILIO|ACUAMAR|UNIBLANC|PURA QUIMICA", solo, re.I) else None


def categoria_de(proveedor, tipo, detalle, pres):
    if tipo == "liquido":
        if ENVASES.match(detalle):
            return "Accesorios"
        return "Líquidos 5L" if pres in ("5 Lts", "10 Lts") else "Líquidos 1L"
    if tipo == "pileta":
        if proveedor.upper().startswith("ACUAMAR") or ENVASES.match(detalle):
            return "Accesorios"
        return "Pileta"
    return "Accesorios"


def numero(valor):
    return valor if isinstance(valor, (int, float)) else None


def lee_excel(ruta):
    ws = openpyxl.load_workbook(ruta, data_only=True)["Lista PRECIO"]
    filas, avisos, vistos = [], [], {}
    proveedor = tipo = ultimo_codigo = None

    cabecera_inicial = limpia(ws.cell(1, 2).value)
    etiqueta = re.search(r"Proveedor:\s*([^|]+)$", cabecera_inicial, re.I)
    if etiqueta:
        proveedor = limpia(etiqueta.group(1))
        tipo = "liquido" if "LIQUIDO" in cabecera_inicial.upper() else "pileta"

    for r in range(3, ws.max_row + 1):
        a, b, c, d, e, f, h = (ws.cell(r, col).value for col in (1, 2, 3, 4, 5, 6, 8))
        if isinstance(a, str) and a.strip().lower() == "codigo":
            cabecera = limpia(b)
            proveedor = proveedor_de(cabecera)
            tipo = "liquido" if "LIQUIDO" in cabecera.upper() else "pileta" if "PILETA" in cabecera.upper() else "otros"
            ultimo_codigo = None
            continue
        detalle = limpia(b)
        if not detalle or not proveedor:
            continue

        codigo = limpia(a)
        if codigo and not re.search(r"\d", codigo) and len(codigo) > 3:
            codigo = ultimo_codigo
        elif codigo:
            ultimo_codigo = codigo
        else:
            codigo = ultimo_codigo

        venta, cant, costo_lista, margen, envase = numero(c), numero(d), numero(e), numero(f), numero(h)
        pres = presentacion(detalle, cant)
        categoria = categoria_de(proveedor, tipo, detalle, pres)
        nota = None
        incompleta = False
        margen_fila = margen
        if venta and venta > 0 and margen is not None and margen > -100:
            costo = redondear(venta / (1 + margen / 100))
            precio = redondear(venta)
            if envase:
                nota = (f"La venta de la lista incluye ${round(envase, 2)} de envase; "
                        f"se guarda el costo equivalente para que la app calcule ese precio.")
        elif venta and venta > 0 and costo_lista and costo_lista > 0:
            costo = redondear(costo_lista)
            precio = redondear(venta)
            margen_fila = redondear((precio / costo - 1) * 100)
            nota = "La lista no trae el margen: se calculó desde la venta y el costo."
        elif venta and venta > 0:
            # La planilla trae la venta pero no el costo: se conserva el costo cargado.
            costo = None
            precio = redondear(venta)
            margen_fila = None
            incompleta = True
            nota = ("La lista no trae el costo de este producto: se mantiene el costo ya cargado "
                    "y su venta de la lista.")
        elif costo_lista and costo_lista > 0:
            costo = redondear(costo_lista)
            precio = redondear(costo * (1 + (margen or 0) / 100))
            nota = "Sin venta en la lista: se calculó desde el costo."
        else:
            avisos.append(f"Fila {r}: {detalle[:50]} no tiene costo ni venta; queda afuera.")
            continue

        fila = {"codigo": codigo or "", "nombre": nombre_corto(detalle, pres), "detalle": detalle,
                "presentacion": pres, "categoria": categoria, "prov": proveedor,
                "costo": costo, "precio": precio,
                "margen": float(margen_fila) if margen_fila is not None else None,
                "alerta": nota, "incompleta": incompleta, "fila": r}
        base_clave = f"{normaliza(fila['codigo']) or normaliza(fila['nombre'])}|{normaliza(pres)}"
        clave = base_clave
        if base_clave in vistos:
            clave = f"{base_clave}|{normaliza(fila['nombre'])}"
            fila["alerta"] = " ".join(x for x in [nota, f"Código y presentación repetidos (fila {vistos[base_clave]}); se distingue por nombre."] if x)
        else:
            vistos[base_clave] = r
        fila["clave"] = clave
        filas.append(fila)
    return filas, avisos


def completar_costos(items, avisos):
    """La planilla no siempre trae el costo (celdas combinadas o en blanco).

    Para esos productos se conserva el costo que ya está en el catálogo publicado:
    el precio de venta sí viene en la lista y es el que manda. Si tampoco hay costo
    anterior, el producto queda afuera con aviso para cargarlo a mano.
    """
    anterior = []
    if os.path.exists(SALIDA_PUBLICADO):
        try:
            with open(SALIDA_PUBLICADO, encoding="utf-8") as f:
                anterior = json.load(f).get("productos", [])
        except (OSError, ValueError):
            anterior = []

    def clave(p):
        return f"{normaliza(p.get('codigo')) or normaliza(p.get('nombre'))}|{normaliza(p.get('presentacion'))}"

    def tamano(p):
        m = re.match(r"^(\d+(?:X\d+)?)", normaliza(p.get("presentacion")))
        return m.group(1) if m else ""

    def tamano_clave(p):
        codigo = normaliza(p.get("codigo"))
        return f"{codigo}|{tamano(p)}" if codigo and tamano(p) else ""

    por_clave, por_codigo, cuentas, por_tamano, cuentas_tamano = {}, {}, {}, {}, {}
    for p in anterior:
        por_clave.setdefault(clave(p), p)
        codigo = normaliza(p.get("codigo"))
        if codigo:
            cuentas[codigo] = cuentas.get(codigo, 0) + 1
            por_codigo[codigo] = p
        tk = tamano_clave(p)
        if tk:
            cuentas_tamano[tk] = cuentas_tamano.get(tk, 0) + 1
            por_tamano[tk] = p

    cuentas_lista, cuentas_lista_tamano = {}, {}
    for it in items:
        codigo = normaliza(it["codigo"])
        if codigo:
            cuentas_lista[codigo] = cuentas_lista.get(codigo, 0) + 1
        tk = tamano_clave(it)
        if tk:
            cuentas_lista_tamano[tk] = cuentas_lista_tamano.get(tk, 0) + 1

    conservados = []
    for it in items:
        if it["costo"] is not None:
            conservados.append(it)
            continue
        previo = por_clave.get(clave(it))
        codigo = normaliza(it["codigo"])
        if not previo and codigo and cuentas_lista.get(codigo) == 1 and cuentas.get(codigo) == 1:
            previo = por_codigo.get(codigo)
        tk = tamano_clave(it)
        if not previo and tk and cuentas_lista_tamano.get(tk) == 1 and cuentas_tamano.get(tk) == 1:
            previo = por_tamano.get(tk)
        if previo and previo.get("costo"):
            it["costo"] = redondear(float(previo["costo"]))
            it["margen"] = redondear((it["precio"] / it["costo"] - 1) * 100)
            it["alerta"] = " ".join(x for x in [it["alerta"], "Se conservó el costo del catálogo anterior."] if x)
            it["incompleta"] = False
            conservados.append(it)
        else:
            avisos.append(f"Fila {it['fila']}: {it['detalle'][:50]} no tiene costo cargado; queda afuera.")
    return conservados, avisos


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else None
    if not ruta:
        sys.exit(__doc__)
    items, avisos = lee_excel(ruta)
    if not items:
        sys.exit("No se encontraron productos en la planilla.")

    items, avisos = completar_costos(items, avisos)

    proveedores = sorted({i["prov"] for i in items})
    prov_id = {nombre: f"pv{n + 1}" for n, nombre in enumerate(proveedores)}
    productos = []
    for it in items:
        productos.append({
            "codigo": it["codigo"], "nombre": it["nombre"], "detalle": it["detalle"],
            "presentacion": it["presentacion"], "categoria": it["categoria"],
            "prov": it["prov"], "provId": prov_id[it["prov"]],
            "costo": it["costo"], "precio": it["precio"], "margen": it["margen"],
            "clave": it["clave"],
            "alerta": it["alerta"],
            # Stock inicial de demostración: la importación nunca toca el stock real.
            "stock": 10, "activo": True,
        })

    cat = {"productos": productos,
           "proveedores": [{"id": prov_id[n], "nombre": n} for n in proveedores]}

    cuerpo = json.dumps(cat, ensure_ascii=False, indent=1)
    cabecera = (
        "// catalogo.js — GENERADO por tools/importar-precios.py. No editar a mano.\n"
        "//\n"
        "// Lista de precios del proveedor convertida al catálogo del negocio.\n"
        "//   costo   = venta ÷ (1 + margen) — la botella ya está incluida en la venta\n"
        "//   precio  = precio de venta de la lista\n"
        "//   margen  = % de ganancia del proveedor\n"
        "//   alerta  = corrección o dato a revisar\n"
        "// Semilla para una instalación nueva. La actualización del negocio en\n"
        "// producción se hace desde Precios → Importar la lista completa del proveedor.\n"
        "const CATALOGO = "
    )
    with open(SALIDA, "w", encoding="utf-8") as f:
        f.write(cabecera + cuerpo + ";\n")

    publicado = {
        "version": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "generado": "tools/importar-precios.py",
        "origen": os.path.basename(ruta),
        "proveedores": cat["proveedores"],
        "productos": productos,
    }
    os.makedirs(os.path.dirname(SALIDA_PUBLICADO), exist_ok=True)
    with open(SALIDA_PUBLICADO, "w", encoding="utf-8") as f:
        json.dump(publicado, f, ensure_ascii=False, indent=1)

    print(f"OK  {len(productos)} productos · {len(proveedores)} proveedores · {len(avisos)} filas descartadas")
    print(f"    -> {SALIDA}")
    print(f"    -> {SALIDA_PUBLICADO}  (version {publicado['version']})")
    for nombre in proveedores:
        n = sum(1 for p in productos if p["prov"] == nombre)
        print(f"       {nombre:<14} {n} productos")
    for aviso in avisos:
        print(f"       [aviso] {aviso}")


if __name__ == "__main__":
    main()
